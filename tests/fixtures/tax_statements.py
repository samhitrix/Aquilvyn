"""Synthetic broker tax statements (no real person's data) shared by unit and end-to-end tests.

* ``zerodha_like_xlsx`` — the shape of Zerodha Console › Reports › Tax P&L: several sheets, lot-level
  "Tradewise Exits" plus per-symbol summaries and totals that repeat the same sales, dividends, charges.
* ``other_broker_csv`` — a different broker's single-table capital-gains CSV (prices instead of values,
  a "Gain Type" column, DD/MM/YYYY dates) to show the reader isn't tied to one broker.
"""
from __future__ import annotations

import io

PAN = "ABCDE1234F"
CLIENT_ID = "ZZ0001"

HEAD = ["Symbol", "ISIN", "Entry Date", "Exit Date", "Quantity", "Buy Value", "Sell Value", "Profit", "Period of Holding",
        "Fair Market Value", "Taxable Profit", "Turnover"]
# (symbol, isin, entry, exit, qty, buy, sell)
INTRADAY = [("NIFTYBEES", "INF204KB14I2", "2026-04-06", "2026-04-06", 4, 1025.32, 1037.04)]
SHORT = [("AUBANK", "INE949L01017", "2025-08-04", "2026-05-14", 4, 2985.40, 3980.00),
         ("KOTAKBANK", "INE237A01036", "2025-08-04", "2026-05-14", 4, 1596.56, 1510.00),
         ("LICI", "INE0J1Y01017", "2025-09-01", "2026-07-28", 15, 0, 6393.00),          # buy value 0: cost basis missing
         ("KARURVYSYA", "INE036D01028", "2025-06-04", "2026-05-14", 8, 1630.16, 1995.70),
         ("KARURVYSYA", "INE036D01028", "2025-08-25", "2026-05-14", 21, 0, 5987.10)]    # 0 again — but an earlier lot shows the cost
LONG = [("PERSISTENT", "INE262H01021", "2023-06-01", "2026-07-27", 10, 30000.00, 61000.00),
        ("ITC", "INE154A01025", "2022-05-10", "2026-06-02", 100, 26000.00, 43000.00),
        ("TATASTEEL", "INE081A01020", "2024-01-15", "2026-05-14", 100, 14000.00, 16000.00)]
MFS = [("AXIS BLUECHIP FUND - DIRECT PLAN", "INF846K01DP8", "2022-04-01", "2026-06-10", 300, 10000.00, 20000.00),   # long
       ("BANDHAN NIFTY 50 INDEX FUND - DIRECT PLAN", "INF194KB1AL4", "2026-01-05", "2026-06-10", 400, 20000.00, 19000.00)]  # short, loss
DIVIDENDS = [("ITC#", "INE154A01025", "2026-05-27", 100, 8, 800.0), ("TATASTEEL", "INE081A01020", "2026-06-12", 100, 4, 400.0)]

EQ_ST = sum(s - b for *_, b, s in SHORT)            # incl. the zero-cost LICI / KARURVYSYA sales
N_SALES = 1 + 5 + 3 + 2                              # intraday + short + long + mutual funds
EQ_LT = sum(s - b for *_, b, s in LONG)             # 50000.00
MF_LT, MF_ST = 10000.0, -1000.0
DIV_TOTAL = sum(d[-1] for d in DIVIDENDS)            # 1200.00


def _lot(r: tuple) -> list:
    sym, isin, e, x, q, b, s = r
    from datetime import date

    days = (date.fromisoformat(x) - date.fromisoformat(e)).days
    return [sym, isin, e, x, q, b, s, round(s - b, 2), days, 0, round(s - b, 2), 0]


def zerodha_like_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    meta = [["Client ID", CLIENT_ID], ["Client Name", "TEST USER"], ["PAN", PAN]]
    ws = wb.active
    ws.title = "Tradewise Exits from 2026-04-01"
    for r in [["View Zerodha's guide on using tax reports for filing."], [], *meta, [], ["Tradewise Exits from 2026-04-01 to 2026-09-15"], [],
              ["Equity - Intraday"], HEAD, *[_lot(r) for r in INTRADAY], [],
              ["Equity - Short Term"], HEAD, *[_lot(r) for r in SHORT], [],
              ["Equity - Long Term"], HEAD, *[_lot(r) for r in LONG], [],
              ["Equity - Buyback"], HEAD, [],
              ["Mutual Funds"], HEAD, *[_lot(r) for r in MFS], [],
              ["F&O"], ["Symbol", "Entry Date", "Exit Date", "Quantity", "Buy Value", "Sell Value", "Profit", "Turnover"]]:
        ws.append(r)
    eq = wb.create_sheet("Equity and Non Equity")
    intraday = round(sum(s - b for *_, b, s in INTRADAY), 2)
    for r in [*meta, [], ["Taxpnl Statement for Equity from 2026-04-01 to 2026-09-15"], ["Realized Profit Breakdown"],
              ["Equity Intraday/Speculative profit", intraday], ["Equity Short Term profit", EQ_ST], ["Equity Long Term profit", EQ_LT], [],
              ["Charges"], ["Account Head", "Amount"], ["Brokerage - Z", 20.0], ["Securities Transaction Tax - Z", 150.0], [],
              ["Equity Short Term"], ["Symbol", "Quantity", "Buy Value", "Sell Value", "Realized P&L"],
              *[[s, q, b, v, round(v - b, 2)] for s, _, _, _, q, b, v in SHORT], ["BANDHAN NIFTY 50 INDEX FUND - DIRECT PLAN", 400, 20000.0, 19000.0, -1000.0], [],
              ["Equity Long Term"], ["Symbol", "Quantity", "Buy Value", "Sell Value", "Realized P&L"],
              *[[s, q, b, v, round(v - b, 2)] for s, _, _, _, q, b, v in LONG], ["AXIS BLUECHIP FUND - DIRECT PLAN", 300, 10000.0, 20000.0, 10000.0]]:
        eq.append(r)
    mf = wb.create_sheet("Mutual Funds")
    for r in [*meta, [], ["Realized Profit Breakdown"], ["Equity Short Term profit", MF_ST], ["Equity Long Term profit", MF_LT]]:
        mf.append(r)
    dv = wb.create_sheet("Dividends & Interest")
    for r in [*meta, [], ["Equity Dividends & Debt interest from 2026-04-01 to 2026-09-15"], ["Total Dividend Amount", DIV_TOTAL],
              ["Total Interest Amount", 0], ["Quarterly Breakdown", "Net Dividend Amount"], ["Upto 15-Jun-2026", DIV_TOTAL], [],
              ["Symbol", "ISIN", "Ex-date", "Quantity", "Dividend Per Share", "Net Dividend Amount"], *[list(d) for d in DIVIDENDS],
              ["Dividends are credited to your registered bank account within 30–45 days from the ex-date."]]:
        dv.append(r)
    od = wb.create_sheet("Other Debits and Credits")
    for r in [*meta, [], ["Equity"], ["Particulars", "Posting Date", "Debit", "Credit"],
              ["Off-market transfer fee for stocks gifted to TEST RELATIVE", "2026-07-27", 88.5, 0]]:
        od.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def other_broker_csv() -> bytes:
    """A different broker's layout: one table, per-unit prices, DD/MM/YYYY dates, an explicit gain-type column."""
    return (
        b"Capital Gains Report,,,,,,,,\n"
        b"Name,TEST USER,,,,,,,\n"
        b"PAN,PQRSX6789K,,,,,,,\n"
        b",,,,,,,,\n"
        b"Stock Name,ISIN,Quantity,Buy Date,Buy Price,Sell Date,Sell Price,Realised P&L,Gain Type\n"
        b"HDFC BANK,INE040A01034,10,05/01/2025,1600,20/06/2026,1750,1500,Long Term\n"
        b"INFOSYS,INE009A01021,5,10/02/2026,1900,15/07/2026,1500,-2000,Short Term\n"
        b"Total,,,,,,,-500,\n"
    )
