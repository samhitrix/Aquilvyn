"""Broker tax statements → realised lots + income (any broker), and the FY tax engine."""
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fixtures import tax_statements as fx  # noqa: E402


@pytest.fixture(scope="module")
def zerodha():
    from portfolio_svc.importers import taxpnl

    return taxpnl.parse(fx.zerodha_like_xlsx())


def test_detects_a_tax_statement_but_not_a_holdings_file():
    from portfolio_svc.importers import detect_kind

    assert detect_kind("taxpnl.xlsx", fx.zerodha_like_xlsx()) == "tax_pnl"
    assert detect_kind("gains.csv", fx.other_broker_csv()) == "tax_pnl"
    holdings = b"Symbol,ISIN,Quantity Available,Average Price\nTCS,INE467B01029,10,3500\n"
    assert detect_kind("holdings.csv", holdings) == "holdings"


def test_zerodha_like_statement_lots_income_and_meta(zerodha):
    rows, issues, meta = zerodha
    gains = [r for r in rows if r.record == "gain"]
    income = [r for r in rows if r.record == "income"]
    # lot-level rows only: the per-symbol summary sheet repeats the same sales and must not be counted again
    assert len(gains) == len(fx.INTRADAY) + len(fx.SHORT) + len(fx.LONG) + len(fx.MFS)
    assert meta["pan"] == fx.PAN and meta["client_id"] == fx.CLIENT_ID and meta["broker"] == "Zerodha"
    assert meta["period"] == ["2026-04-01", "2026-09-15"] and meta["charges"] == 170.0
    by = {(r.symbol, r.term, r.asset) for r in gains}
    assert ("NIFTYBEES", "intraday", "equity") in by and ("PERSISTENT", "long", "equity") in by
    assert ("AXIS BLUECHIP FUND - DIRECT PLAN", "long", "equity_mf") in by            # MF term from the holding period
    assert ("BANDHAN NIFTY 50 INDEX FUND - DIRECT PLAN", "short", "equity_mf") in by
    assert sorted((r.symbol, r.profit) for r in income) == [("ITC", Decimal("800")), ("TATASTEEL", Decimal("400"))]  # '#' stripped
    assert all(r.term == "dividend" for r in income)


def test_self_checks_reconcile_totals_and_flag_problems(zerodha):
    rows, issues, meta = zerodha
    assert meta["checks"] and all(c["ok"] for c in meta["checks"]), [c for c in meta["checks"] if not c["ok"]]
    labels = " | ".join(c["label"] for c in meta["checks"])
    assert "Equity long term profit" in labels and "Mutual funds · Equity long term profit" in labels and "Total dividend amount" in labels
    lici = next(r for r in rows if r.symbol == "LICI")
    assert "zero_cost" in lici.flags
    text = " ".join(i["error"] for i in issues)
    assert "buy value of 0 (KARURVYSYA, LICI)" in text and "gifted to TEST RELATIVE" in text


def test_a_broken_statement_is_caught_by_the_reconciliation():
    import io

    from openpyxl import load_workbook

    from portfolio_svc.importers import taxpnl

    wb = load_workbook(io.BytesIO(fx.zerodha_like_xlsx()))
    ws = wb["Tradewise Exits from 2026-04-01"]
    for row in ws.iter_rows():
        if row[0].value == "PERSISTENT":
            row[10].value = 1.0  # taxable profit corrupted
            row[7].value = 1.0
    buf = io.BytesIO()
    wb.save(buf)
    _, issues, meta = taxpnl.parse(buf.getvalue())
    assert any(not c["ok"] and "long term" in c["label"].lower() for c in meta["checks"])
    assert any("statement says" in i["error"] for i in issues)


def test_other_broker_layout_prices_gain_type_and_totals_row():
    from portfolio_svc.importers import taxpnl

    rows, issues, meta = taxpnl.parse(fx.other_broker_csv())
    assert meta["pan"] == "PQRSX6789K" and "broker" not in meta
    assert [(r.symbol, r.term, r.buy_date, r.sell_date, r.buy_value, r.sell_value, r.profit) for r in rows] == [
        ("HDFC BANK", "long", date(2025, 1, 5), date(2026, 6, 20), Decimal("16000"), Decimal("17500"), Decimal("1500")),
        ("INFOSYS", "short", date(2026, 2, 10), date(2026, 7, 15), Decimal("9500"), Decimal("7500"), Decimal("-2000")),
    ]  # the "Total" row is not a sale


def test_tax_engine_buckets_exemption_setoff_and_cess(zerodha):
    from portfolio_svc import tax

    rows, _, _ = zerodha
    gains = [{"asset": r.asset, "term": r.term, "sell_date": r.sell_date, "taxable_profit": r.taxable_profit, "sell_value": r.sell_value,
              "cost_override": None} for r in rows if r.record == "gain"]
    inc = [{"kind": r.term, "amount": r.profit} for r in rows if r.record == "income"]
    s = tax.compute("2026-27", gains, inc, 30, today=date(2026, 9, 29))
    line = {x["key"]: x for x in s["lines"]}
    st = fx.EQ_ST + fx.MF_ST                       # 6301.04 (MF short-term loss nets within the bucket)
    lt = fx.EQ_LT + fx.MF_LT                       # 60000 → fully inside the ₹1.25 L exemption
    assert line["st_equity"]["taxable"] == pytest.approx(st) and line["st_equity"]["tax"] == pytest.approx(st * 0.20, abs=0.01)
    assert line["lt_equity"]["gross"] == pytest.approx(lt) and line["lt_equity"]["tax"] == 0
    assert s["ltcg_exemption"] == {"limit": 125000.0, "used": lt, "left": 125000.0 - lt}
    assert line["dividends"]["tax"] == pytest.approx(fx.DIV_TOTAL * 0.30, abs=0.01)
    base = st * 0.20 + fx.DIV_TOTAL * 0.30 + line["intraday"]["tax"]
    assert s["estimated_tax"] == pytest.approx(round(base * 1.04, 2), abs=0.02)
    assert s["advance_tax"]["due_date"] == "2026-12-15" and s["advance_tax"]["cumulative_pct"] == 75


def test_losses_set_off_in_the_right_order_and_carry_forward():
    from portfolio_svc import tax

    d = date(2026, 6, 1)

    def g(asset, term, p):
        return {"asset": asset, "term": term, "sell_date": d, "taxable_profit": p, "sell_value": 0, "cost_override": None}
    # short-term loss offsets short-term gains first, then long-term; long-term loss only long-term
    s = tax.compute("2026-27", [g("equity", "short", -50_000), g("equity", "short", 20_000), g("equity", "long", 200_000),
                                g("equity", "long", -10_000)], [], 30)
    lt = next(x for x in s["lines"] if x["key"] == "lt_equity")
    assert lt["taxable"] == pytest.approx(200_000 - 10_000 - 30_000 - 125_000)
    s = tax.compute("2026-27", [g("equity", "long", -40_000), g("equity", "short", 10_000)], [], 30)
    assert s["carry_forward"]["long_term_loss"] == 40_000 and next(x for x in s["lines"] if x["key"] == "st_equity")["taxable"] == 10_000
    s = tax.compute("2026-27", [g("equity", "intraday", -5_000), g("equity", "short", 10_000)], [], 30)
    assert s["carry_forward"]["speculative_loss"] == 5_000  # intraday losses don't touch capital gains


def test_entered_cost_replaces_a_zero_buy_value():
    from portfolio_svc import tax

    lot = {"asset": "equity", "term": "short", "sell_date": date(2026, 7, 28), "taxable_profit": Decimal("6393"), "sell_value": Decimal("6393"),
           "cost_override": Decimal("5000")}
    assert tax.effective_gain(lot) == 1393.0


def test_old_rates_before_23_july_2024():
    from portfolio_svc import tax

    s = tax.compute("2024-25", [{"asset": "equity", "term": "short", "sell_date": date(2024, 6, 1), "taxable_profit": 10_000,
                                 "sell_value": 0, "cost_override": None}], [], 30)
    assert next(x for x in s["lines"] if x["key"] == "st_equity")["tax"] == 1500.0


def _lt(taxable_lt=0.0, taxable_st=0.0):
    from portfolio_svc import tax

    d = date(2026, 6, 1)
    g = [{"asset": "equity", "term": "long", "sell_date": d, "taxable_profit": taxable_lt, "sell_value": 0, "cost_override": None},
         {"asset": "equity", "term": "short", "sell_date": d, "taxable_profit": taxable_st, "sell_value": 0, "cost_override": None}]
    return tax.compute("2026-27", g, [], 30)


HOLD = [
    {"instrument_id": "i-itc", "asset_type": "stock", "symbol": "ITC.NS", "name": "ITC", "price": 400, "lots": [{"buy_date": "2024-01-01", "qty": 100, "cost": 300}]},
    {"instrument_id": "i-axis", "asset_type": "mutual_fund", "symbol": "1", "name": "Axis Bluechip", "price": 60, "meta": {"mf_category": "equity:large_cap"},
     "lots": [{"buy_date": "2023-01-01", "qty": 1000, "cost": 40}]},
    {"instrument_id": "i-liq", "asset_type": "mutual_fund", "symbol": "2", "name": "Liquid Fund", "price": 10, "meta": {"mf_category": "debt:liquid"},
     "lots": [{"buy_date": "2023-01-01", "qty": 1000, "cost": 8}]},
    {"instrument_id": "i-yes", "asset_type": "stock", "symbol": "YESBANK.NS", "name": "Yes Bank", "price": 20, "lots": [{"buy_date": "2026-05-01", "qty": 1000, "cost": 25}]},
]


def test_gain_harvesting_fills_the_exemption_with_exact_units_and_skips_debt():
    from portfolio_svc import tax

    h = tax.harvest("2026-27", _lt(taxable_lt=100_000), HOLD, date(2026, 9, 29))   # 25,000 of the 1.25 L left
    gains = {a["symbol"]: a for a in h["actions"] if a["kind"] == "gain"}
    assert set(gains) == {"1", "ITC.NS"}                                   # the liquid (debt) fund is never suggested
    assert gains["1"]["quantity"] == 1000 and gains["1"]["booked"] == 20_000   # biggest first
    assert gains["ITC.NS"]["quantity"] == 50 and gains["ITC.NS"]["booked"] == 5_000 and h["room_left"] == 0


def test_loss_harvesting_says_units_and_the_real_tax_saved():
    from portfolio_svc import tax

    h = tax.harvest("2026-27", _lt(taxable_st=50_000), HOLD, date(2026, 9, 29))
    yes = next(a for a in h["actions"] if a["kind"] == "loss")
    assert (yes["symbol"], yes["quantity"], yes["booked"]) == ("YESBANK.NS", 1000, -5_000)
    assert yes["tax_saved"] == round(5_000 * 0.20 * 1.04, 2) and h["total_saving"] == yes["tax_saved"]
    # nothing to offset → the loss is still listed, but saves nothing now
    h0 = tax.harvest("2026-27", _lt(), HOLD, date(2026, 9, 29))
    assert next(a for a in h0["actions"] if a["kind"] == "loss")["tax_saved"] == 0


def test_fifo_sells_oldest_units_first_and_done_actions_are_not_repeated():
    from portfolio_svc import tax

    row = {"instrument_id": "i-x", "asset_type": "stock", "symbol": "X", "name": "X", "price": 100,
           "lots": [{"buy_date": "2026-02-01", "qty": 10, "cost": 90},    # +100 on the oldest units
                    {"buy_date": "2026-03-01", "qty": 40, "cost": 120}]}  # −800 on the newer ones
    h = tax.harvest("2026-27", _lt(taxable_st=10_000), [row], date(2026, 9, 29))
    x = h["actions"][0]
    assert x["quantity"] == 50 and x["booked"] == -700      # can't sell only the loss-making lot: FIFO takes the old ones first
    done = [{"key": "loss:i-x", "kind": "loss", "st_part": -700, "lt_part": 0, "reflected": False}]
    h2 = tax.harvest("2026-27", _lt(taxable_st=10_000), [row], date(2026, 9, 29), done=done)
    assert h2["actions"] == [] and h2["taxable_left"]["st_eq"] == 9_300   # marked done: gone, and its offset is already counted
