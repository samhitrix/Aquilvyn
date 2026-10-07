"""Mutual-fund portfolio disclosures → stock-level holdings per scheme (E16 look-through).

Every fund house (AMC) must publish each scheme's full portfolio monthly (SEBI). They do it as Excel
workbooks — usually one sheet per scheme, a few title rows (AMC, scheme name, "portfolio as on …") above
a table of instrument name / ISIN / industry or rating / quantity / market value / "% to NAV" — sometimes
zipped, sometimes one file per scheme. Layouts differ by AMC, so nothing here is keyed to an AMC: the
table is found by its columns, the scheme by its title rows (or the sheet name), and the scheme is
matched to AMFI's scheme names. This module is pure (bytes in, data out) so every layout can be tested
offline; fetching lives in ``fund_sources.py``.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from typing import Any

XLSX_MAGIC = ZIP_MAGIC = b"PK\x03\x04"  # an .xlsx is itself a zip
XLS_MAGIC = b"\xd0\xcf\x11\xe0"
ISIN_RE = re.compile(r"^[A-Z]{2}[0-9A-Z]{9}[0-9]$")
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


# ---------- reading any spreadsheet (xlsx / xls / csv / zip of those) ----------

def _xlsx(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _xls(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    out = []
    for sh in book.sheets():
        rows = []
        for i in range(sh.nrows):
            row = []
            for cell in sh.row(i):
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    except Exception:  # noqa: BLE001 — keep the raw number
                        pass
                row.append(v)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _csv(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [("csv", [list(r) for r in csv.reader(io.StringIO(text), dialect)])]


def read_tables(content: bytes, filename: str = "") -> list[tuple[str, list[list[Any]]]]:
    """(sheet name, rows) for every sheet; a .zip is opened and every spreadsheet inside read."""
    name = filename.lower()
    if content[:4] == ZIP_MAGIC and not (name.endswith(".xlsx") or name.endswith(".xlsm")):
        try:
            zf = zipfile.ZipFile(io.BytesIO(content))
            inner = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xls", ".xlsm", ".csv")) and not n.startswith("__MACOSX")]
            if inner and "[Content_Types].xml" not in zf.namelist():
                out: list[tuple[str, list[list[Any]]]] = []
                for n in inner:
                    for sheet, rows in read_tables(zf.read(n), n):
                        out.append((f"{n.rsplit('/', 1)[-1]}:{sheet}" if len(inner) > 1 else sheet, rows))
                return out
        except zipfile.BadZipFile:
            pass
    if content[:4] == XLSX_MAGIC:
        return _xlsx(content)
    if content[:4] == XLS_MAGIC:
        return _xls(content)
    return _csv(content)


# ---------- finding the holdings table ----------

def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\n", " ")).strip().lower()


def _is_pct_header(h: str) -> bool:
    if "weight" in h:
        return True
    pct = "%" in h or "percent" in h or "perc" in h
    return pct and any(k in h for k in ("nav", "net asset", "aum", "asset", "portfolio", "total"))


def _is_name_header(h: str) -> bool:
    return any(k in h for k in ("name", "instrument", "issuer", "company", "security", "scrip", "particular")) and "isin" not in h


def header_columns(row: list[Any]) -> dict[str, int] | None:
    """Column indexes of a holdings table header, or None if this row isn't one."""
    hs = [_norm(c) for c in row]
    isin = next((i for i, h in enumerate(hs) if h in ("isin", "isin code", "isin no", "isin no.", "isin number") or h.startswith("isin")), None)
    pct = next((i for i, h in enumerate(hs) if _is_pct_header(h)), None)
    if isin is None or pct is None:
        return None
    name = next((i for i, h in enumerate(hs) if _is_name_header(h)), None)
    industry = next((i for i, h in enumerate(hs) if any(k in h for k in ("industry", "sector", "rating"))), None)
    return {"isin": isin, "pct": pct, "name": name if name is not None else -1, "industry": industry if industry is not None else -1}


def _num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    s = str(v or "").strip().replace(",", "").replace("%", "").replace("$", "")
    if not s or s in ("-", "--", "nil", "NIL"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    try:
        x = float(s.strip("()"))
    except ValueError:
        return None
    return -x if neg else x


def isin_kind(isin: str, industry: str = "") -> str:
    """equity (Indian shares) | foreign_equity | fund (units of another fund / ETF) | debt."""
    if isin.startswith("INF"):
        return "fund"
    if isin.startswith("INE") or isin.startswith("IN9"):
        # Indian ISIN: IN + E + 4-char issuer + 2-char security type ("01" = equity shares) + serial + check digit
        return "equity" if isin[7:9] in ("01", "02") else "debt"
    if isin.startswith("IN"):
        return "debt"  # IN0… / IN1… / IN2… = government securities, SDLs, T-bills
    ind = industry.lower()
    return "debt" if any(k in ind for k in ("aaa", "aa+", "sovereign", "treasury", "bond")) else "foreign_equity"


# ---------- scheme name and date from the title rows ----------

_PLAN_WORDS = {"direct", "regular", "plan", "growth", "option", "idcw", "dividend", "payout", "reinvestment", "reinvest", "bonus",
               "the", "scheme", "of", "an", "a", "erstwhile", "formerly", "known", "as"}


def clean_scheme_title(text: str) -> str:
    """"Portfolio Statement of SBI Small Cap Fund as on 31/08/2026 (An open ended …)" → "SBI Small Cap Fund"."""
    t = re.sub(r"\(.*?\)", " ", str(text))  # "(An open ended dynamic equity scheme …)"
    t = re.sub(r"(?i)^\s*(scheme\s*name|name\s*of\s*(the\s*)?scheme|(monthly\s*)?(statement\s*of\s*)?portfolio\s*(statement)?\s*(of|for)?|"
               r"monthly\s*portfolio\s*(statement)?\s*(of|for)?)\s*[:\-–]?\s*", "", t)
    t = re.sub(r"(?i)\s*(portfolio|monthly portfolio statement|statement)\s*(as\s*(on|at)|for the month).*$", "", t)
    t = re.sub(r"(?i)[\s,:\-–]*\b(as\s*(on|at|of)|for\s*the\s*(month|period)|month\s*end(ed)?)\b.*$", "", t)
    return re.sub(r"\s+", " ", t).strip(" :-–,")


def _title_candidates(rows: list[list[Any]]) -> list[str]:
    out: list[str] = []  # explicit "Scheme name:" labels
    plain: list[str] = []
    for r in rows:
        cells = [str(c).strip() for c in r if c not in (None, "") and not isinstance(c, int | float | datetime | date)]
        if not cells:
            continue
        joined = " ".join(cells)
        low = joined.lower()
        if "scheme name" in low or "name of the scheme" in low or "name of scheme" in low:
            # "SCHEME NAME :" | "HDFC Flexi Cap Fund (…)"  — or one cell "Scheme Name: HDFC Flexi Cap Fund"
            label_idx = next((i for i, c in enumerate(cells) if re.search(r"(?i)scheme", c)), 0)
            after = cells[label_idx + 1:]
            out.insert(0, clean_scheme_title(" ".join(after) if after else joined))
            continue
        for c in cells:
            cl = c.lower()
            if " fund" in cl or cl.endswith("fund") or " etf" in cl or "bees" in cl or "fof" in cl:
                plain.append(clean_scheme_title(c))
    # nearest the table first; the AMC banner ("X Mutual Fund", "… monthly portfolio …") last
    plain.reverse()
    plain.sort(key=lambda c: ("mutual fund" in c.lower() or "portfolio" in c.lower()))
    return [c for c in out + plain if len(c) >= 6]


def find_as_of(rows: list[list[Any]]) -> date | None:
    for r in rows:
        for c in r:
            if isinstance(c, datetime):
                return c.date()
            if isinstance(c, date):
                return c
        text = " ".join(str(c) for c in r if c not in (None, ""))
        d = parse_date_text(text)
        if d:
            return d
    return None


def parse_date_text(text: str) -> date | None:
    """"31st July 2026", "31-Jul-26", "July 31,2026", "31.07.2026", "2026-07-31", "July 2026" (with "as on" / "month")."""
    t = re.sub(r"(?<=\d)(?=[a-z])|(?<=[a-z])(?=\d)", " ", text.lower().replace("_", " "))  # "31jul2026" → "31 jul 2026"
    m = re.search(r"(?<!\d)(\d{1,2})(?:st|nd|rd|th)?[\s\-./,']+([a-z]{3,9})[\s\-./,']+(\d{4}|\d{2})(?!\d)", t)
    if m and m.group(2)[:3] in MONTHS:
        return _safe_date(_year(m.group(3)), MONTHS[m.group(2)[:3]], int(m.group(1)))
    m = re.search(r"\b([a-z]{3,9})[\s\-./]+(\d{1,2})(?:st|nd|rd|th)?[\s,\-./]+(\d{4})(?!\d)", t)
    if m and m.group(1)[:3] in MONTHS:
        return _safe_date(int(m.group(3)), MONTHS[m.group(1)[:3]], int(m.group(2)))
    m = re.search(r"(?<!\d)(\d{1,2})[./\-](\d{1,2})[./\-](\d{4})(?!\d)", t)
    if m:
        return _safe_date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.search(r"(?<!\d)(20\d{2})[./\-](\d{1,2})[./\-](\d{1,2})(?!\d)", t)
    if m:
        return _safe_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"\b([a-z]{3,9})[\s\-',.]+(\d{4})(?!\d)", t)  # "July 2026" → that month's last day
    if m and m.group(1)[:3] in MONTHS and ("as on" in t or "as of" in t or "portfolio" in t or "month" in t):
        return month_end(int(m.group(2)), MONTHS[m.group(1)[:3]])
    return None


def month_end(y: int, mo: int) -> date | None:
    if not 2000 <= y <= 2100:
        return None
    return (date(y + (mo == 12), mo % 12 + 1, 1) - timedelta(days=1))


def _year(s: str) -> int:
    return int(s) + 2000 if len(s) == 2 else int(s)


def _safe_date(y: int, mo: int, d: int) -> date | None:
    try:
        return date(y, mo, d) if 2000 <= y <= 2100 else None
    except ValueError:
        return None


# ---------- one workbook → schemes ----------

def parse_portfolio(content: bytes, filename: str = "") -> list[dict[str, Any]]:
    """Every scheme found in a portfolio disclosure file:
    [{sheet, scheme, as_of, holdings: [{isin, name, industry, kind, weight}], equity_pct, total_pct}]"""
    out: list[dict[str, Any]] = []
    for sheet, rows in read_tables(content, filename):
        out.extend(_parse_sheet(sheet, rows))
    return out


def _parse_sheet(sheet: str, rows: list[list[Any]]) -> list[dict[str, Any]]:
    """A sheet normally holds one scheme; some AMCs stack several schemes in one sheet, each with its own
    title rows and header — every header found starts a new scheme."""
    headers = [(i, cols) for i, r in enumerate(rows[:5000]) if (cols := header_columns(r))]
    if not headers:
        return []
    schemes = []
    for n, (hi, cols) in enumerate(headers):
        end = headers[n + 1][0] if n + 1 < len(headers) else len(rows)
        title_from = headers[n - 1][0] + 1 if n else 0
        title_rows = rows[max(title_from, hi - 12):hi]
        holdings = _rows_to_holdings(rows[hi + 1:end], cols)
        if not holdings:
            continue
        names = _title_candidates(title_rows)
        schemes.append({
            "sheet": sheet, "scheme": names[0] if names else clean_scheme_title(sheet), "titles": names[:4],
            "as_of": (find_as_of(title_rows) or find_as_of(rows[:hi])),
            "holdings": holdings,
            "equity_pct": round(sum(h["weight"] for h in holdings if h["kind"] in ("equity", "foreign_equity")), 2),
            "total_pct": round(sum(h["weight"] for h in holdings), 2),
        })
    return schemes


def _rows_to_holdings(rows: list[list[Any]], cols: dict[str, int]) -> list[dict[str, Any]]:
    raw = []
    for r in rows:
        def cell(key: str, row: list[Any] = r) -> Any:
            i = cols[key]
            return row[i] if 0 <= i < len(row) else None
        isin = str(cell("isin") or "").strip().upper().replace(" ", "")
        if not ISIN_RE.match(isin):
            continue
        w = _num(cell("pct"))
        if w is None:
            continue
        name = str(cell("name") or "").strip() or isin
        industry = str(cell("industry") or "").strip()
        raw.append({"isin": isin, "name": name, "industry": industry, "kind": isin_kind(isin, industry), "weight": w})
    if not raw:
        return []
    # "% to NAV" is a percentage in most files and a fraction (0.0912) in some: a full portfolio sums to ~100 or ~1
    total = sum(h["weight"] for h in raw)
    if 0 < total <= 1.5:
        for h in raw:
            h["weight"] *= 100
    # the same ISIN can appear twice (e.g. listed + locked-in shares) → one line
    merged: dict[str, dict[str, Any]] = {}
    for h in raw:
        m = merged.setdefault(h["isin"], {**h, "weight": 0.0})
        m["weight"] += h["weight"]
    return [{**h, "weight": round(h["weight"], 4)} for h in merged.values() if h["weight"] > 0]


# ---------- matching a sheet to an AMFI scheme ----------

_JOINED = re.compile(r"\b(small|mid|large|flexi|multi|micro)(cap)\b")
_GENERIC_BRAND = ("mutual", "fund", "asset", "management", "company", "limited", "ltd", "amc", "india", "investment", "investments")


def _words(name: str) -> list[str]:
    t = re.sub(r"\(.*?\)", " ", name.lower()).replace("&", " and ")
    t = _JOINED.sub(r"\1 \2", re.sub(r"[^a-z0-9 ]+", " ", t))  # "Midcap" = "Mid Cap"
    return [w for w in t.split() if w not in _PLAN_WORDS and w not in ("mutual", "and")]


def _brand(amc: str) -> set[str]:
    return {w for w in re.sub(r"[^a-z0-9 ]+", " ", amc.lower()).split() if w not in _GENERIC_BRAND}


def scheme_tokens(name: str, amc: str = "") -> set[str]:
    brand = _brand(amc)
    return {w for w in _words(name) if w not in brand}


def acronyms(name: str, amc: str = "") -> set[str]:
    """Short codes fund houses name sheets / files with: "Mirae Asset Mid Cap Fund" → mamcf; "SBI Small Cap Fund" → sscf, sbiscf."""
    words, brand = _words(name), _brand(amc)
    if not words:
        return set()
    out = {"".join(w[0] for w in words), "".join(w if w in brand else w[0] for w in words),
           "".join(w[0] for w in words if w not in brand)}
    if words[-1] == "fund":  # "…cf" and "…c" both seen
        out |= {a[:-1] for a in out}
    return {a for a in out if len(a) >= 4}


def code_of(name: str) -> str:
    """A sheet / file name as a bare code: "MAMCF_Aug2026.xlsx" → "mamcf", "SSCF" → "sscf", "x.zip:SSCF" → "sscf"."""
    stem = name.lower().rsplit("/", 1)[-1].rsplit(":", 1)[-1]
    stem = re.sub(r"\?.*$", "", stem)
    stem = re.sub(r"\.(xlsx|xls|xlsm|csv|zip)$", "", stem)
    stem = re.sub(r"([\s_\-]+\d{1,2}(st|nd|rd|th)?)?[\s_\-]*(" + "|".join(MONTHS) + r")[a-z]*[\s_\-]*\d{2,4}.*$", "", stem)
    return re.sub(r"[^a-z0-9]", "", stem)


def match_score(sheet_name: str, amfi_name: str, amc: str = "") -> float:
    a, b = scheme_tokens(sheet_name, amc), scheme_tokens(amfi_name, amc)
    if not a or not b:
        return 0.0
    jac = len(a & b) / len(a | b)
    seq = SequenceMatcher(None, " ".join(sorted(a)), " ".join(sorted(b))).ratio()
    score = max(jac, 0.6 * jac + 0.4 * seq)
    if b <= a and len(a - b) <= 2:  # every word of the fund's name is in the title, plus a word or two
        score = max(score, 0.75 + 0.25 * jac)
    return round(score, 3)


def match_schemes(schemes: list[dict[str, Any]], targets: dict[str, dict[str, str]], threshold: float = 0.72,
                  file_name: str = "") -> dict[str, dict[str, Any]]:
    """targets: {amfi_code: {name, amc}} → {amfi_code: {index, score, scheme}} for the best sheet above the threshold.
    A sheet (or a one-scheme file) named with the fund's initials ("MAMCF") counts as a strong match."""
    codes = [{code_of(s.get("sheet", ""))} for s in schemes]
    if len(schemes) == 1 and file_name:  # one scheme per file: the file name is its code
        codes[0].add(code_of(file_name))
    out: dict[str, dict[str, Any]] = {}
    for code, t in targets.items():
        acr = acronyms(t["name"], t.get("amc", ""))
        best: tuple[float, int] | None = None
        for i, s in enumerate(schemes):
            score = max([match_score(n, t["name"], t.get("amc", "")) for n in (s.get("titles") or [s["scheme"]])] + [0.0])
            if acr & codes[i] and sum(1 for c in codes if acr & c) == 1:
                score = max(score, 0.9)
            if best is None or score > best[0]:
                best = (score, i)
        if best and best[0] >= threshold:
            out[code] = {"index": best[1], "score": best[0], "scheme": schemes[best[1]]["scheme"]}
    return out


def closest(schemes: list[dict[str, Any]], target: dict[str, str], n: int = 3) -> list[tuple[str, float]]:
    """For a 'not found' report: the scheme titles in a file nearest to the fund we looked for."""
    scored = []
    for s in schemes:
        best = max(((match_score(t, target["name"], target.get("amc", "")), t) for t in (s.get("titles") or [s["scheme"]])), default=(0.0, s["scheme"]))
        scored.append((f"{best[1]} [sheet {s.get('sheet', '')}]", best[0]))
    return sorted(scored, key=lambda x: -x[1])[:n]
