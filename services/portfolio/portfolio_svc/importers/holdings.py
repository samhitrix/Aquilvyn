"""Holdings statements (a snapshot of what you own *now*), as opposed to transaction files.

Supported (auto-detected by column names, any sheet):
* Zerodha Console → Portfolio → Holdings → Download (XLSX: "Equity" + "Mutual Funds" sheets)
* Kite web → Holdings → Download CSV ("Instrument, Qty., Avg. cost, LTP, …")
* Generic broker/RTA holdings CSV/XLSX with a name/symbol/ISIN, a quantity and an average cost

Each row becomes one (estimated) opening position at its average cost. A snapshot gives the
correct *current* picture — quantity, invested, value, gain — even when older trades are
missing; only purchase dates are approximate (so XIRR/tax lots are estimates).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

from .base import dec

COLS: dict[str, tuple[str, ...]] = {
    "symbol": ("symbol", "tradingsymbol", "trading symbol", "stock symbol", "instrument", "instrument name", "scrip", "scrip name",
               "script name", "script", "stock", "stock name", "security", "security description",
               "security name", "scheme", "scheme name", "fund", "fund name", "name", "company", "company name"),
    "isin": ("isin", "isin code", "isin no", "isin number"),
    "qty": ("quantity available", "qty.", "qty", "quantity", "units", "balance units", "closing units", "closing balance",
            "total quantity", "holding quantity", "free quantity", "net quantity", "qty available", "available qty", "holding qty",
            "current qty", "balance qty", "demat qty"),
    "qty_discrepant": ("quantity discrepant",),
    "qty_long_term": ("quantity long term",),
    "qty_pledged_margin": ("quantity pledged (margin)",),
    "qty_pledged_loan": ("quantity pledged (loan)",),
    "avg": ("average price", "avg. cost", "avg cost", "average cost", "avg price", "average buy price", "buy average", "buy avg",
            "avg. price", "purchase nav", "avg nav", "average nav", "cost price", "avg cost price", "average cost price", "avg. buy price",
            "avg buy price", "average buy price", "buy avg price", "avg rate", "average rate", "wac", "avg. cost price"),
    "invested": ("invested value", "invested", "investment", "invested amount", "cost value", "total cost", "purchase value", "amount invested",
                 "value at cost", "buy value", "investment value", "holding cost", "total investment"),
    "price": ("previous closing price", "ltp", "last price", "current price", "close price", "closing price", "market price",
              "nav", "current nav", "latest nav", "current market price", "cmp", "market rate", "closing rate"),
    "pnl": ("unrealized p&l", "unrealised p&l", "p&l", "unrealized pnl", "unrealised pnl"),
    "itype": ("instrument type", "asset class", "asset type", "type", "category"),
    "sector": ("sector",),
}


@dataclass(slots=True)
class HoldingRow:
    name: str
    symbol: str | None          # exchange symbol for stocks/ETFs; None for mutual funds (resolved via ISIN/name)
    isin: str | None
    asset_type: str             # stock | etf | mutual_fund
    quantity: Decimal
    avg_price: Decimal
    invested: Decimal
    price: Decimal | None       # price printed on the statement (previous close / NAV) — fallback only
    long_term_qty: Decimal = Decimal(0)
    source_row: str = ""
    sector: str | None = None   # as printed on the statement (Zerodha has a Sector column)
    price_date: date | None = None  # date the statement's price is for (NPS: "NAV as on …"); default: import day


def _norm(h: Any) -> str:
    """Lower-case, single spaces, currency marks dropped ("Avg Price (₹)" → "avg price")."""
    s = re.sub(r"\((?:₹|rs\.?|inr|in ₹|in rs\.?)\)|₹", "", str(h or "").lower())
    return re.sub(r"\s+", " ", s).strip(" :*")


def _header_map(row: list[Any]) -> dict[str, int] | None:
    cells = [_norm(c) for c in row]
    found: dict[str, int] = {}
    for key, names in COLS.items():
        for n in names:  # names are in priority order
            if n in cells:
                found[key] = cells.index(n)
                break
    if "qty" in found and ("avg" in found or "invested" in found) and ("symbol" in found or "isin" in found):
        return found
    return None


def _tables_xlsx(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    # Not read_only: Zerodha's export declares its sheet size as just "A1", and openpyxl's
    # read-only mode trusts that — it would only ever see column A ("no holdings table found").
    wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _tables_xls(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    """Legacy Excel 97-2003 (.xls)."""
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    return [(sh.name, [sh.row_values(i) for i in range(sh.nrows)]) for sh in book.sheets()]


XLSX_MAGIC, XLS_MAGIC = b"PK\x03\x04", b"\xd0\xcf\x11\xe0"


def _tables(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    if content[:4] == XLSX_MAGIC:
        return _tables_xlsx(content)
    if content[:4] == XLS_MAGIC:
        return _tables_xls(content)
    return _tables_csv(content)


def _tables_csv(content: bytes) -> list[tuple[str, list[list[Any]]]]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [("csv", [list(r) for r in csv.reader(io.StringIO(text), dialect)])]


def looks_like_holdings(content: bytes) -> bool:
    try:
        tables = _tables(content)
    except Exception:
        return False
    return any(_header_map(r) for _, rows in tables for r in rows[:60])


def _cell(row: list[Any], idx: int | None) -> Any:
    return row[idx] if idx is not None and idx < len(row) else None


def _num(v: Any) -> Decimal:
    if isinstance(v, int | float):
        return Decimal(str(v))
    return dec(v)


def _kind(sheet: str, name: str, symbol_col_is_scheme: bool, isin: str | None, itype: str, sector: str = "") -> str:
    s, t, sec = sheet.lower(), itype.lower(), sector.strip().lower()
    ticker_like = " " not in name.strip()  # exchange symbols have no spaces; fund names do
    if "mutual" in s or "mutual" in t or t in ("mf", "fund") or symbol_col_is_scheme:
        return "mutual_fund"
    scheme_words = re.search(r"\b(fund|plan|growth|idcw|dividend|fof|reinvest|payout)\b", name, re.I)
    etf_words = re.search(r"\b(etf|bees)\b|bees$", name, re.I)
    if isin and isin.startswith("INF") and not ticker_like and (scheme_words or not etf_words):
        # a scheme name — even when its category says "Index Funds/ETFs" (Zerodha's MF rows); an exchange-traded
        # fund's name ("NIPPON INDIA ETF NIFTY 50 BEES") has no scheme words and is an ETF
        return "mutual_fund"
    if sec == "etf" or (ticker_like and ("etf" in t or name.upper().split("-")[0].endswith("BEES"))) or " ETF" in f" {name.upper()}":
        return "etf"
    return "stock"


def _sector(raw: str, kind: str) -> str | None:
    v = raw.strip()
    if kind != "stock" or not v or v in ("-", "NA", "N/A"):
        return None
    return v.title() if v.isupper() else v


def _exchange_symbol(name: str) -> str | None:
    sym = re.sub(r"[^A-Z0-9&\-_.]", "", name.upper())
    # Zerodha appends a series/unit suffix: "LIQUIDBEES-F" (fractional units), "GOLDBEES-E",
    # "-BE"/"-BZ" (trade-for-trade series)… The exchange ticker is the part before it.
    # (Real tickers with a hyphen, like BAJAJ-AUTO, have longer suffixes and are kept.)
    sym = re.sub(r"-(?:E|F|BE|BZ|BL|SM|ST|IL|IV|RR|GS|N\d?)$", "", sym)
    return sym or None


CLIENT_ID_LABELS = ("client id", "client code", "ucc", "client id:")


def statement_meta(content: bytes) -> dict[str, str]:
    """Account identity printed above the table (Zerodha Console: 'Client ID | AB1234'), used to
    route the next statement from the same account to the same family member."""
    try:
        tables = _tables(content)
    except Exception:
        return {}
    for _, rows in tables:
        for r in rows[:25]:
            cells = [str(c).strip() for c in r if c not in (None, "")]
            for i, c in enumerate(cells[:-1]):
                if _norm(c).rstrip(":") in CLIENT_ID_LABELS and re.fullmatch(r"[A-Za-z0-9]{4,12}", cells[i + 1]):
                    return {"client_id": cells[i + 1].upper()}
    return {}


def parse(content: bytes) -> tuple[list[HoldingRow], list[dict[str, Any]]]:
    return parse_tables(_tables(content))


def parse_tables(tables: list[tuple[str, list[list[Any]]]]) -> tuple[list[HoldingRow], list[dict[str, Any]]]:
    # Zerodha's export has one sheet per instrument type (Equity, Mutual Funds) *and* a Combined
    # sheet with everything. Read only Combined when it has a holdings table, else every sheet.
    combined = [(n, rows) for n, rows in tables if "combined" in n.lower() and any(_header_map(r) for r in rows[:80])]
    if combined:
        tables = combined
    rows_out: list[HoldingRow] = []
    errors: list[dict[str, Any]] = []
    for sheet, rows in tables:
        hdr: dict[str, int] | None = None
        hdr_cells: list[str] = []
        for r in rows:
            again = _header_map(r)
            if again:  # a (new) table starts here — e.g. Equity, then Mutual Funds, in one sheet
                hdr, hdr_cells = again, [_norm(c) for c in r]
                continue
            if hdr is None:
                continue
            name = str(_cell(r, hdr.get("symbol")) or "").strip()
            isin = str(_cell(r, hdr.get("isin")) or "").strip().upper() or None
            if not name and not isin:
                continue
            if name.lower() in ("total", "grand total", "net total"):
                continue
            label = f"{sheet} · {name or isin}"
            try:
                qty = _num(_cell(r, hdr["qty"]))
                disc = _num(_cell(r, hdr.get("qty_discrepant")))
                pm, pl = _num(_cell(r, hdr.get("qty_pledged_margin"))), _num(_cell(r, hdr.get("qty_pledged_loan")))
                avg = _num(_cell(r, hdr.get("avg")))
                invested = _num(_cell(r, hdr.get("invested")))
                price_raw = _cell(r, hdr.get("price"))
                price = _num(price_raw) if price_raw not in (None, "") else None
                pnl_raw = _cell(r, hdr.get("pnl"))
                # Zerodha splits quantity into available / discrepant / pledged; pick the total that
                # reproduces the statement's own unrealised P&L (qty × (price − avg)).
                candidates = [qty, qty + disc, qty + disc + pm + pl]
                if pnl_raw not in (None, "") and price is not None and avg and price != avg:
                    implied = _num(pnl_raw) / (price - avg)
                    qty = min(candidates, key=lambda c: abs(c - implied))
                    # double-check against the statement's own P&L: a mismatch means we misread the row
                    if abs(qty - implied) > max(Decimal("0.01"), abs(implied) * Decimal("0.02")):
                        errors.append({"row": label, "level": "warning",
                                       "error": f"Quantity {qty} doesn't match this row's Unrealized P&L (implies ≈{implied:.3f}) — please check it in the review."})
                else:
                    qty = qty + disc
                if qty <= 0:
                    continue
                if not avg and invested:
                    avg = (invested / qty).quantize(Decimal("0.0001"))
                if not invested:
                    invested = (qty * avg).quantize(Decimal("0.01"))
                if avg <= 0:
                    errors.append({"row": label, "error": "No average cost on this row — add it as a transaction instead"})
                    continue
                itype = str(_cell(r, hdr.get("itype")) or "")
                scheme_col = hdr_cells[hdr["symbol"]] in ("scheme", "scheme name", "fund", "fund name") if "symbol" in hdr else False
                kind = _kind(sheet, name, scheme_col, isin, itype, str(_cell(r, hdr.get("sector")) or ""))
                sym = None if kind == "mutual_fund" else _exchange_symbol(name)
                rows_out.append(HoldingRow(
                    name=name or (isin or ""), symbol=sym, isin=isin, asset_type=kind, quantity=qty, avg_price=avg,
                    invested=invested, price=price, long_term_qty=min(_num(_cell(r, hdr.get("qty_long_term"))), qty), source_row=label,
                    sector=_sector(str(_cell(r, hdr.get("sector")) or ""), kind),
                ))
            except Exception as exc:
                errors.append({"row": label, "error": f"Could not read this row: {exc}"})
    if not rows_out and not errors:
        errors.append({"row": "file", "error": "No holdings table found. Expected columns like Symbol/ISIN, Quantity and Average price."})
    rows_out, repeats = _dedupe(rows_out)
    if repeats:
        errors.append({"row": "file", "level": "warning",
                       "error": f"{repeats} row(s) appear more than once in the file (the same holding on another sheet/table) — each holding is counted once."})
    return rows_out, errors


def _dedupe(rows: list[HoldingRow]) -> tuple[list[HoldingRow], int]:
    """Brokers' Excel exports can list the same holdings on several sheets (e.g. 'Equity' and a
    combined view). The same instrument with the same quantity and average price is one holding,
    not two. (Different rows of one instrument — e.g. LIQUIDBEES and LIQUIDBEES-F fractional
    units — have different names/quantities and are still added together later.)"""
    seen: set[tuple[str, str, Decimal, Decimal]] = set()
    out: list[HoldingRow] = []
    for r in rows:
        key = (r.isin or "", r.name.strip().upper(), r.quantity, r.avg_price)
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out, len(rows) - len(out)
